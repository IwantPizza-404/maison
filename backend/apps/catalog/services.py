"""Импорт и экспорт товаров в Excel.

Один формат для всего: экспорт, шаблон и импорт — это один и тот же лист «Products»,
поэтому выгруженный файл можно поправить и загрузить обратно.

- Товар ищется по `slug`: найден — обновляется, пусто — создаётся новый.
- Колонки, которых нет в файле, не трогаются: можно загрузить только `slug` и `price`.
- Импорт всё-или-ничего: при любой ошибке не сохраняется ни одна строка.
- Справочники (бренд, категория, материалы…) ищутся по slug или по названию на любом языке.
- Фото — ссылки в колонке images. Файлы скачиваются при проверке, а записываются
  в хранилище только при подтверждённом импорте без ошибок.
"""

import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import TYPE_CHECKING, Any

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import models, transaction
from django.db.models import QuerySet
from django.utils import translation
from django.utils.text import slugify
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Brand, Category, Product, ProductImage, SpecOption, WatchSpecs
from .photos import Photo, PhotoError, PhotoFetcher, absolute_url

if TYPE_CHECKING:
    from django_stubs_ext import StrOrPromise

SHEET = "Products"
MAX_ROWS = 2000
MAX_PHOTOS_PER_PRODUCT = 15
MAX_PHOTO_DOWNLOADS = 300
LANGS = [code for code, _name in settings.LANGUAGES]
YES, NO = "yes", "no"
TRUE_VALUES = {"yes", "y", "да", "ha", "true", "1", "+", "✓"}
FALSE_VALUES = {"no", "n", "нет", "yo‘q", "yo'q", "yoq", "false", "0", "-", ""}


@dataclass(frozen=True)
class Column:
    key: str
    kind: str  # text, decimal, int, bool, choice, brand, category, option, options, images
    target: str = "product"  # product | specs | images
    required: bool = False
    choices: type[models.TextChoices] | None = None
    option_kind: str = ""
    help: "StrOrPromise" = ""
    width: int = 14

    @property
    def is_relation(self) -> bool:
        return self.kind in {"brand", "category", "option"}

    def label(self) -> str:
        """Название колонки на языке админки — из verbose_name поля модели."""
        if self.kind == "images":
            return str(ProductImage._meta.verbose_name_plural)
        model = WatchSpecs if self.target == "specs" else Product
        base, _sep, lang = self.key.rpartition("_")
        if lang in LANGS and base in {"name", "description"}:
            return f"{field_label(model, base)} ({lang})"
        return field_label(model, self.key)


def field_label(model: type[models.Model], name: str) -> str:
    model_field = model._meta.get_field(name)
    assert isinstance(model_field, models.Field | models.ManyToManyField)
    return str(model_field.verbose_name)


C = Column
COLUMNS: list[Column] = [
    C("slug", "text", width=28, help=gettext_lazy("Empty — a new product is created.")),
    C("brand", "brand", required=True, width=18),
    C("category", "category", required=True),
    C("name_en", "text", required=True, width=28),
    C("name_ru", "text", width=28),
    C("name_uz", "text", width=28),
    C("reference", "text"),
    C("price", "decimal", required=True, help=gettext_lazy("USD, number without currency.")),
    C("compare_at_price", "decimal"),
    C("availability", "choice", choices=Product.Availability),
    C("stock_quantity", "int"),
    C("lead_time_days", "int"),
    C("condition", "choice", choices=Product.Condition),
    C("gender", "choice", choices=Product.Gender),
    C("is_published", "bool", help=gettext_lazy("New products are drafts unless “yes”.")),
    C("is_featured", "bool"),
    C("description_en", "text", width=40),
    C("description_ru", "text", width=40),
    C("description_uz", "text", width=40),
    C(
        "images",
        "images",
        "images",
        width=60,
        help=gettext_lazy(
            "Photo links, one per line; the first is the main photo. "
            "Empty — photos stay as they are."
        ),
    ),
    C("movement", "choice", "specs", choices=WatchSpecs.Movement),
    C("caliber", "text", "specs"),
    C("power_reserve_hours", "int", "specs"),
    C("case_material", "option", "specs", option_kind=SpecOption.Kind.CASE_MATERIAL),
    C("case_diameter_mm", "decimal", "specs"),
    C("case_thickness_mm", "decimal", "specs"),
    C("glass", "option", "specs", option_kind=SpecOption.Kind.GLASS),
    C("water_resistance_m", "int", "specs"),
    C("dial_color", "option", "specs", option_kind=SpecOption.Kind.DIAL_COLOR),
    C("strap_material", "option", "specs", option_kind=SpecOption.Kind.STRAP_MATERIAL),
    C(
        "complications",
        "options",
        "specs",
        option_kind=SpecOption.Kind.COMPLICATION,
        width=28,
        help=gettext_lazy("Several values separated by commas."),
    ),
    C("year", "int", "specs"),
    C("has_box", "bool", "specs"),
    C("has_papers", "bool", "specs"),
    C("serial_number", "text", "specs"),
]
COLUMNS_BY_KEY = {column.key: column for column in COLUMNS}


class ImportFileError(Exception):
    """Файл нельзя прочитать как таблицу товаров."""


# --- Справочники ----------------------------------------------------------------


class Lookups:
    """Справочники в памяти: один запрос на весь файл, поиск по slug и названиям."""

    def __init__(self) -> None:
        self.brands = list(Brand.objects.all())
        self.categories = list(Category.objects.all())
        self.options: dict[str, list[SpecOption]] = {}
        for option in SpecOption.objects.all():
            self.options.setdefault(option.kind, []).append(option)

    @staticmethod
    def _find(items: list, value: str, names: list[str], what: str) -> Any:
        needle = value.strip().casefold()
        by_slug = [item for item in items if item.slug == needle]
        if by_slug:
            return by_slug[0]
        found = {
            item.pk: item
            for item in items
            for name in names
            if str(getattr(item, name, "") or "").casefold() == needle
        }
        if len(found) > 1:
            raise ValueError(
                _("“%(value)s” matches several %(what)s.") % {"value": value, "what": what}
            )
        if not found:
            raise ValueError(_("%(what)s “%(value)s” not found.") % {"value": value, "what": what})
        return next(iter(found.values()))

    def brand(self, value: str) -> Brand:
        return self._find(self.brands, value, ["name"], str(Brand._meta.verbose_name))

    def category(self, value: str) -> Category:
        names = [f"name_{lang}" for lang in LANGS]
        return self._find(self.categories, value, names, str(Category._meta.verbose_name))

    def option(self, kind: str, value: str) -> SpecOption:
        names = [f"name_{lang}" for lang in LANGS]
        what = str(SpecOption.Kind(kind).label).lower()
        return self._find(self.options.get(kind, []), value, names, what)


# --- Экспорт и шаблон -----------------------------------------------------------


def cell_value(column: Column, product: Product, specs: WatchSpecs | None) -> Any:
    source = specs if column.target == "specs" else product
    if source is None:
        return None
    if column.kind == "images":
        return "\n".join(absolute_url(image.image) for image in product.images.all()) or None
    if column.kind == "options":
        return ", ".join(option.slug for option in source.complications.all())  # type: ignore[union-attr]
    value = getattr(source, column.key)
    if column.kind == "bool":
        return YES if value else NO
    if column.kind == "brand":
        return value.name
    if column.kind in {"category", "option"}:
        return value.slug if value else None
    return value if value != "" else None


def build_workbook(products: list[Product]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET
    lists = workbook.create_sheet("Lists")
    lists.sheet_state = "hidden"

    header_fill = PatternFill("solid", fgColor="1F2937")
    required_fill = PatternFill("solid", fgColor="92400E")
    for index, column in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=column.key)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill if column.required else header_fill
        note = column.label() + (f"\n{column.help}" if column.help else "")
        cell.comment = Comment(note, "Maison", width=240, height=80)
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.freeze_panes = "B2"

    for row, product in enumerate(products, start=2):
        specs = getattr(product, "watch_specs", None)
        for index, column in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=index, value=cell_value(column, product, specs))
            if column.kind == "images":
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    add_dropdowns(sheet, lists, last_row=max(len(products) + 500, 1000))
    add_instructions(workbook)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def add_dropdowns(sheet, lists, last_row: int) -> None:
    """Выпадающие списки допустимых значений — из скрытого листа Lists."""
    lookups = Lookups()
    sources: dict[str, list[str]] = {
        "brand": [brand.name for brand in lookups.brands],
        "category": [category.slug for category in lookups.categories],
        "bool": [YES, NO],
    }
    list_column = 0
    for index, column in enumerate(COLUMNS, start=1):
        if column.kind == "choice" and column.choices:
            values = list(column.choices.values)
        elif column.kind == "option":
            values = [option.slug for option in lookups.options.get(column.option_kind, [])]
        elif column.kind in sources:
            values = sources[column.kind]
        else:
            continue
        if not values:
            continue
        list_column += 1
        letter = get_column_letter(list_column)
        for row, value in enumerate(values, start=1):
            lists.cell(row=row, column=list_column, value=value)
        validation = DataValidation(
            type="list", formula1=f"=Lists!${letter}$1:${letter}${len(values)}", allow_blank=True
        )
        validation.error = _("Choose a value from the list.")
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        target = get_column_letter(index)
        validation.add(f"{target}2:{target}{last_row}")


def add_instructions(workbook) -> None:
    sheet = workbook.create_sheet(_("Instructions"))
    rows = [
        [_("How to fill in")],
        [_("One row — one product. Do not rename the columns of the “Products” sheet.")],
        [_("Existing slug — the product is updated; empty slug — a new product is created.")],
        [_("Columns removed from the file are left unchanged in the catalog.")],
        [_("If any row has an error, nothing is imported — fix the file and upload it again.")],
        [_("Photos: links in the “images” column — Google Drive, Dropbox or any direct link.")],
        [],
        [_("Column"), _("Field"), _("Required"), _("Values")],
    ]
    for column in COLUMNS:
        if column.choices:
            values = ", ".join(f"{value} — {label}" for value, label in column.choices.choices)
        elif column.kind == "bool":
            values = f"{YES} / {NO}"
        elif column.kind in {"brand", "category", "option", "options"}:
            values = _("from the list, slug or name")
        else:
            values = str(column.help)
        required = _("yes") if column.required else ""
        rows.append([column.key, column.label(), required, values])
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(bold=True, size=14)
    for cell in sheet[8]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCD", (24, 32, 12, 80), strict=True):
        sheet.column_dimensions[letter].width = width
    for row in sheet.iter_rows(min_row=9):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")


def export_products(queryset: QuerySet[Product]) -> bytes:
    products = list(
        queryset.select_related(
            "brand",
            "category",
            "watch_specs__case_material",
            "watch_specs__glass",
            "watch_specs__dial_color",
            "watch_specs__strap_material",
        )
        .prefetch_related("watch_specs__complications", "images")
        .order_by("brand__name", "name")
    )
    return build_workbook(products)


def products_template() -> bytes:
    return build_workbook([])


# --- Импорт ---------------------------------------------------------------------


@dataclass
class RowError:
    row: int
    column: str
    message: str


@dataclass
class RowResult:
    row: int
    action: str  # create | update | unchanged
    title: str
    changes: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    rows: list[RowResult] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def count(self, action: str) -> int:
        return sum(1 for row in self.rows if row.action == action)

    @property
    def created(self) -> int:
        return self.count("create")

    @property
    def updated(self) -> int:
        return self.count("update")

    @property
    def unchanged(self) -> int:
        return self.count("unchanged")


def read_rows(data: bytes) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(_("Could not read the file. Upload an .xlsx file.")) from exc
    sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value).strip().lower() if value is not None else "" for value in next(rows, ())]
    if "name_en" not in header and "slug" not in header:
        raise ImportFileError(_("The first row must contain column names, as in the template."))
    unknown = [key for key in header if key and key not in COLUMNS_BY_KEY]
    if unknown:
        raise ImportFileError(_("Unknown columns: %s") % ", ".join(unknown))
    duplicates = {key for key in header if key and header.count(key) > 1}
    if duplicates:
        raise ImportFileError(_("Duplicate columns: %s") % ", ".join(sorted(duplicates)))

    result = []
    for number, values in enumerate(rows, start=2):
        row = {key: value for key, value in zip(header, values, strict=False) if key}
        if all(value is None or str(value).strip() == "" for value in row.values()):
            continue
        result.append((number, row))
    workbook.close()
    if len(result) > MAX_ROWS:
        raise ImportFileError(_("Too many rows: at most %d per file.") % MAX_ROWS)
    return [key for key in header if key], result


def as_text(value: Any) -> str:
    if value is None:
        return ""
    # Excel хранит артикул 116610 как число 116610.0
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def choice_value(choices: type[models.TextChoices], text: str) -> str:
    needle = text.casefold()
    if needle in choices.values:
        return needle
    for lang in LANGS:
        with translation.override(lang):
            for value, label in choices.choices:
                if str(label).casefold() == needle:
                    return value
    allowed = ", ".join(choices.values)
    raise ValueError(
        _("Unknown value “%(value)s”. Allowed: %(allowed)s.") % {"value": text, "allowed": allowed}
    )


def parse(column: Column, raw: Any, lookups: Lookups) -> Any:
    """Значение ячейки → значение поля. Пустая ячейка — пустое значение поля."""
    text = as_text(raw)
    if column.kind == "bool":
        if text.casefold() in TRUE_VALUES:
            return True
        if text.casefold() in FALSE_VALUES:
            return False
        raise ValueError(_("Enter yes or no."))
    if column.kind == "options":
        parts = [part.strip() for part in text.replace(";", ",").split(",") if part.strip()]
        return [lookups.option(column.option_kind, part) for part in parts]
    if not text:
        return "" if column.kind == "text" else None
    if column.kind == "text":
        return text
    if column.kind in {"decimal", "int"}:
        try:
            number = Decimal(text.replace(" ", "").replace(" ", "").replace(",", "."))
        except InvalidOperation as exc:
            raise ValueError(_("Enter a number.")) from exc
        if column.kind == "decimal":
            return number
        if number != number.to_integral_value() or number < 0:
            raise ValueError(_("Enter a whole number."))
        return int(number)
    if column.kind == "choice":
        assert column.choices is not None
        return choice_value(column.choices, text)
    if column.kind == "brand":
        return lookups.brand(text)
    if column.kind == "category":
        return lookups.category(text)
    if column.kind == "option":
        return lookups.option(column.option_kind, text)
    raise AssertionError(column.kind)


def unique_slug(product: Product, taken: set[str] | None = None) -> str:
    """Slug из бренда и английского названия; при совпадении — с номером: rolex-gmt-2."""
    taken = taken or set()
    base = slugify(f"{product.brand.name} {product.name_en}")[:200] or "product"  # type: ignore[attr-defined]
    slug, number = base, 2
    while slug in taken or Product.objects.filter(slug=slug).exists():
        slug, number = f"{base}-{number}", number + 1
    return slug


def split_urls(text: str) -> list[str]:
    """Ссылки через перенос строки, пробел или «;» — без повторов, порядок сохраняется."""
    urls = text.replace(";", " ").split()
    return list(dict.fromkeys(urls))


def error_column(field_name: str) -> str:
    # modeltranslation проверяет поле name — в файле это колонка name_en
    if field_name in {"name", "description"}:
        return f"{field_name}_en"
    return "" if field_name == "__all__" else field_name


PhotoPlan = list[ProductImage | Photo]


class RowImporter:
    def __init__(
        self, header: list[str], lookups: Lookups, report: ImportReport, photos: PhotoFetcher
    ) -> None:
        self.header = header
        self.lookups = lookups
        self.report = report
        self.photos = photos
        self.slugs_in_file: set[str] = set()
        # Фото пишутся в хранилище только в конце успешного импорта — см. save_photos
        self.pending_photos: list[tuple[Product, PhotoPlan]] = []

    def error(self, row: int, column: str, message: str) -> None:
        self.report.errors.append(RowError(row, column, message))

    def run(self, number: int, row: dict[str, Any]) -> None:
        errors_before = len(self.report.errors)
        slug = as_text(row.get("slug"))
        product = Product.objects.filter(slug=slug).first() if slug else None
        creating = product is None
        if product is None:
            product = Product(slug=slug)
        specs: WatchSpecs | None = getattr(product, "watch_specs", None) if not creating else None

        if slug:
            if slug in self.slugs_in_file:
                self.error(
                    number, "slug", _("This slug is already used in another row of the file.")
                )
            self.slugs_in_file.add(slug)

        changes: list[str] = []
        complications: list[SpecOption] | None = None
        for key in self.header:
            column = COLUMNS_BY_KEY[key]
            if key == "slug" or column.kind == "images":
                continue
            try:
                value = parse(column, row.get(key), self.lookups)
            except ValueError as exc:
                self.error(number, key, str(exc))
                continue
            if column.required and value in (None, ""):
                self.error(number, key, _("Required column."))
                continue
            if column.kind == "options":
                old = {o.pk for o in specs.complications.all()} if specs and specs.pk else set()
                if old != {o.pk for o in value}:
                    changes.append(key)
                complications = value
                if value and specs is None:
                    specs = WatchSpecs()
                continue
            if column.target == "specs":
                if specs is None:
                    if value in (None, "", False):
                        continue  # у товара нет характеристик и в файле они пустые
                    specs = WatchSpecs()
                target: models.Model = specs
            else:
                target = product
            # Переводы (name_ru…) пустые как NULL — иначе modeltranslation не подставит английский
            if value == "" and target._meta.get_field(key).null:
                value = None
            # Связи сравниваем по id: у нового товара product.brand ещё не задан и бросает ошибку
            if column.is_relation:
                old_value = getattr(target, f"{key}_id")
                new_value = value.pk if value is not None else None
            else:
                old_value, new_value = getattr(target, key), value
            # Пустая строка и NULL — одно и то же: админка хранит пустые описания как ""
            if old_value != new_value and not (old_value in ("", None) and new_value in ("", None)):
                changes.append(key)
                setattr(target, key, value)

        if creating:
            for column in COLUMNS:
                if column.required and column.key not in self.header:
                    self.error(number, column.key, _("Required column is missing in the file."))
        if len(self.report.errors) > errors_before:
            return

        if creating and not product.slug:
            product.slug = unique_slug(product, self.slugs_in_file)
            self.slugs_in_file.add(product.slug)
        if not self.validate(number, product, specs):
            return
        photo_plan: PhotoPlan | None = None
        if "images" in self.header:
            photo_plan, photos_ok = self.plan_photos(number, product, row.get("images"))
            if not photos_ok:
                return
            if photo_plan is not None:
                changes.append("images")

        title = f"{product.brand.name} {product.name_en}"  # type: ignore[attr-defined]
        if not creating and not changes:
            self.report.rows.append(RowResult(number, "unchanged", title))
            return
        product.save()
        if specs is not None:
            specs.product = product
            specs.save()
            if complications is not None:
                specs.complications.set(complications)
        if photo_plan is not None:
            self.pending_photos.append((product, photo_plan))
        action = "create" if creating else "update"
        self.report.rows.append(RowResult(number, action, title, [] if creating else changes))

    def plan_photos(self, number: int, product: Product, raw: Any) -> tuple[PhotoPlan | None, bool]:
        """Новый список фото товара. None — фото не меняются (пустая ячейка или те же ссылки)."""
        urls = split_urls(as_text(raw))
        if not urls:
            return None, True
        if len(urls) > MAX_PHOTOS_PER_PRODUCT:
            self.error(
                number, "images", _("At most %d photos per product.") % MAX_PHOTOS_PER_PRODUCT
            )
            return None, False
        current = (
            {absolute_url(image.image): image for image in product.images.all()}
            if product.pk
            else {}
        )
        if list(current) == urls:
            return None, True
        plan: PhotoPlan = []
        ok = True
        for position, url in enumerate(urls, start=1):
            if url in current:
                plan.append(current[url])
                continue
            try:
                plan.append(self.photos.get(url))
            except PhotoError as exc:
                self.error(
                    number, "images", _("Photo %(n)d: %(error)s.") % {"n": position, "error": exc}
                )
                ok = False
        return (plan if ok else None), ok

    def save_photos(self) -> None:
        for product, plan in self.pending_photos:
            kept = [item.pk for item in plan if isinstance(item, ProductImage)]
            for image in product.images.exclude(pk__in=kept):
                image.delete()  # файлы удалятся после коммита (сигнал)
            for position, item in enumerate(plan):
                if isinstance(item, ProductImage):
                    if item.position != position:
                        item.position = position
                        item.save(update_fields=["position"])
                    continue
                image = ProductImage(product=product, position=position)
                name = f"{product.slug}-{position + 1}.{item.extension}"
                image.image.save(name, ContentFile(item.content), save=False)
                image.save()

    def validate(self, number: int, product: Product, specs: WatchSpecs | None) -> bool:
        valid = True
        checks: list[tuple[models.Model, list[str]]] = [(product, [])]
        if specs is not None:
            checks.append((specs, ["product"]))
        for instance, exclude in checks:
            try:
                instance.full_clean(exclude=exclude)
            except ValidationError as exc:
                valid = False
                for field_name, messages in exc.message_dict.items():
                    self.error(number, error_column(field_name), " ".join(messages))
        return valid


def import_products(data: bytes, *, commit: bool) -> ImportReport:
    """Проверяет файл и, если ошибок нет и commit=True, сохраняет товары.

    Проверка и сохранение — один и тот же код в транзакции: при проверке она откатывается,
    поэтому предпросмотр в точности совпадает с тем, что будет сохранено.
    """
    header, rows = read_rows(data)
    report = ImportReport()
    if not rows:
        raise ImportFileError(_("The file has no products."))
    photos = PhotoFetcher(own_photo_files() if "images" in header else {})
    if "images" in header:
        prefetch_photos(photos, rows)
    with transaction.atomic():
        importer = RowImporter(header, Lookups(), report, photos)
        for number, row in rows:
            importer.run(number, row)
        if commit and report.ok:
            importer.save_photos()
        else:
            transaction.set_rollback(True)
    return report


def own_photo_files() -> dict[str, str]:
    """Фото, уже лежащие на сайте: ссылка из экспорта → имя файла в хранилище."""
    return {
        absolute_url(image.image): str(image.image.name) for image in ProductImage.objects.all()
    }


def prefetch_photos(photos: PhotoFetcher, rows: list[tuple[int, dict[str, Any]]]) -> None:
    """Внешние ссылки скачиваются заранее и параллельно, чтобы импорт уложился в таймаут."""
    external = {
        url
        for _number, row in rows
        for url in split_urls(as_text(row.get("images")))
        if url not in photos.own_files
    }
    if len(external) > MAX_PHOTO_DOWNLOADS:
        raise ImportFileError(
            _("Too many new photos: at most %d per file. Split the file into parts.")
            % MAX_PHOTO_DOWNLOADS
        )
    photos.prefetch(external)
