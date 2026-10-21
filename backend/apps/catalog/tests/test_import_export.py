import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.catalog.models import Product, ProductImage, SpecOption, WatchSpecs
from apps.catalog.services import (
    COLUMNS,
    ImportFileError,
    export_products,
    import_products,
    products_template,
)
from apps.catalog.tests.factories import (
    BrandFactory,
    CategoryFactory,
    ProductFactory,
    SpecOptionFactory,
    WatchSpecsFactory,
)

pytestmark = pytest.mark.django_db


def xlsx(*rows: list) -> bytes:
    """Файл импорта: первая строка — заголовки."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def refs():
    return {
        "brand": BrandFactory(name="Rolex", slug="rolex"),
        "category": CategoryFactory(name="Watches", slug="watches"),
        "steel": SpecOptionFactory(
            kind=SpecOption.Kind.CASE_MATERIAL, slug="steel", name_en="Steel", name_ru="Сталь"
        ),
        "date": SpecOptionFactory(kind=SpecOption.Kind.COMPLICATION, slug="date", name_en="Date"),
        "gmt": SpecOptionFactory(kind=SpecOption.Kind.COMPLICATION, slug="gmt", name_en="GMT"),
    }


HEADER = ["slug", "brand", "category", "name_en", "price", "stock_quantity", "case_material"]


def test_template_has_all_columns_and_dropdowns(refs):
    workbook = load_workbook(io.BytesIO(products_template()))

    header = [cell.value for cell in workbook["Products"][1]]
    assert header == [column.key for column in COLUMNS]
    assert workbook["Products"].max_row == 1
    assert workbook["Lists"].sheet_state == "hidden"
    assert workbook["Products"].data_validations.dataValidation


def test_import_creates_product_with_specs(refs):
    data = xlsx(HEADER, ["", "rolex", "Watches", "Submariner Date", "12 500,50", 2, "Сталь"])

    report = import_products(data, commit=True)

    assert report.ok, report.errors
    assert report.created == 1
    product = Product.objects.get(name_en="Submariner Date")
    assert product.slug == "rolex-submariner-date"
    assert product.brand == refs["brand"]
    assert product.price == Decimal("12500.50")
    assert product.is_published is False  # новые товары — черновики
    assert product.watch_specs.case_material == refs["steel"]


def test_check_does_not_save(refs):
    data = xlsx(HEADER, ["", "rolex", "watches", "Daytona", 30000, 1, ""])

    report = import_products(data, commit=False)

    assert report.ok
    assert report.created == 1
    assert not Product.objects.filter(name_en="Daytona").exists()


def test_partial_update_keeps_other_columns(refs):
    product = ProductFactory(slug="sub", brand=refs["brand"], reference="126610", price=100)

    report = import_products(xlsx(["slug", "price"], ["sub", 150]), commit=True)

    assert report.updated == 1
    assert report.rows[0].changes == ["price"]
    product.refresh_from_db()
    assert product.price == 150
    assert product.reference == "126610"


def test_any_error_imports_nothing(refs):
    data = xlsx(
        HEADER,
        ["", "rolex", "watches", "Good row", 100, 1, ""],
        ["", "Omega", "watches", "Unknown brand", 100, 1, ""],
        ["", "rolex", "watches", "Bad price", "abc", 1, ""],
    )

    report = import_products(data, commit=True)

    assert not report.ok
    assert [(e.row, e.column) for e in report.errors] == [(3, "brand"), (4, "price")]
    assert Product.objects.count() == 0


def test_model_rules_are_checked(refs):
    # В наличии без количества — то же правило, что и в админке
    report = import_products(
        xlsx(HEADER, ["", "rolex", "watches", "No stock", 100, 0, ""]), commit=False
    )

    assert [(e.row, e.column) for e in report.errors] == [(2, "stock_quantity")]


def test_new_product_needs_required_columns(refs):
    report = import_products(xlsx(["name_en", "price"], ["Lonely", 100]), commit=False)

    assert {e.column for e in report.errors} == {"brand", "category"}


def test_duplicate_slug_in_file(refs):
    data = xlsx(
        ["slug", "brand", "category", "name_en", "price", "stock_quantity"],
        ["same", "rolex", "watches", "A", 100, 1],
        ["same", "rolex", "watches", "B", 100, 1],
    )

    report = import_products(data, commit=False)

    assert (3, "slug") in [(e.row, e.column) for e in report.errors]


def test_choices_by_code_or_label_and_booleans(refs):
    header = [
        "brand",
        "category",
        "name_en",
        "price",
        "availability",
        "lead_time_days",
        "is_published",
    ]
    report = import_products(
        xlsx(header, ["rolex", "watches", "On order watch", 100, "On order", 30, "да"]), commit=True
    )

    assert report.ok, report.errors
    product = Product.objects.get(name_en="On order watch")
    assert product.availability == Product.Availability.ON_ORDER
    assert product.is_published is True


def test_complications_list(refs):
    header = ["brand", "category", "name_en", "price", "stock_quantity", "complications"]

    import_products(xlsx(header, ["rolex", "watches", "GMT", 100, 1, "date; GMT"]), commit=True)

    specs = WatchSpecs.objects.get(product__name_en="GMT")
    assert set(specs.complications.values_list("slug", flat=True)) == {"date", "gmt"}


def test_export_then_import_changes_nothing(refs):
    specs = WatchSpecsFactory(
        product=ProductFactory(brand=refs["brand"], category=refs["category"], reference="116610"),
        case_material=refs["steel"],
        has_box=True,
    )
    specs.complications.set([refs["date"]])

    report = import_products(export_products(Product.objects.all()), commit=False)

    assert report.ok, report.errors
    assert report.unchanged == 1


def test_empty_string_and_null_are_the_same(refs):
    # Товар из админки хранит пустое описание как "", импорт пишет NULL — это не изменение
    ProductFactory(slug="admin-made", description_ru="")

    report = import_products(xlsx(["slug", "description_ru"], ["admin-made", None]), commit=False)

    assert report.unchanged == 1


def test_unknown_column_rejects_file(refs):
    with pytest.raises(ImportFileError):
        import_products(xlsx(["slug", "colour"], ["x", "red"]), commit=False)


def test_not_an_excel_file():
    with pytest.raises(ImportFileError):
        import_products(b"not a spreadsheet", commit=False)


# --- Админка ---------------------------------------------------------------------


@pytest.fixture
def admin_client(client, django_user_model):
    client.force_login(django_user_model.objects.create_superuser("a", "a@maison.uz", "p"))
    return client


def test_admin_downloads(admin_client, refs):
    ProductFactory(brand=refs["brand"])
    for name in ("catalog_product_export_all", "catalog_product_download_template"):
        response = admin_client.get(reverse(f"admin:{name}"))

        assert response.status_code == 200
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")
        load_workbook(io.BytesIO(response.content))


def test_admin_export_selected(admin_client, refs):
    keep, _other = ProductFactory(brand=refs["brand"]), ProductFactory(brand=refs["brand"])

    response = admin_client.post(
        reverse("admin:catalog_product_changelist"),
        {"action": "export_selected", "_selected_action": [keep.pk]},
    )

    rows = list(load_workbook(io.BytesIO(response.content))["Products"].values)
    assert [row[0] for row in rows[1:]] == [keep.slug]


def test_admin_import_check_then_confirm(admin_client, refs):
    url = reverse("admin:catalog_product_import_file")
    upload = io.BytesIO(xlsx(HEADER, ["", "rolex", "watches", "Explorer", 7000, 1, ""]))
    upload.name = "products.xlsx"

    checked = admin_client.post(url, {"file": upload})

    assert checked.status_code == 200
    assert not Product.objects.filter(name_en="Explorer").exists()
    token = checked.context["token"]

    confirmed = admin_client.post(url, {"token": token})

    assert confirmed.status_code == 302
    assert Product.objects.filter(name_en="Explorer").exists()


def test_admin_import_shows_errors(admin_client, refs):
    upload = io.BytesIO(xlsx(HEADER, ["", "nope", "watches", "X", 1, 1, ""]))
    upload.name = "products.xlsx"

    response = admin_client.post(reverse("admin:catalog_product_import_file"), {"file": upload})

    assert "token" not in response.context
    assert response.context["report"].errors[0].column == "brand"


def test_admin_builds_slug_from_brand_and_name(admin_client, refs):
    ProductFactory(slug="rolex-explorer")
    url = reverse("admin:catalog_product_add")
    data = {
        "brand": refs["brand"].pk,
        "category": refs["category"].pk,
        "name_en": "Explorer",
        "slug": "",
        "price": "100",
        "availability": "in_stock",
        "stock_quantity": "1",
        "condition": "new",
        "gender": "unisex",
        "watch_specs-TOTAL_FORMS": "0",
        "watch_specs-INITIAL_FORMS": "0",
        "images-TOTAL_FORMS": "0",
        "images-INITIAL_FORMS": "0",
    }

    response = admin_client.post(url, data)

    assert response.status_code == 302, response.context["adminform"].form.errors
    assert Product.objects.filter(slug="rolex-explorer-2", name_en="Explorer").exists()


# --- Фото ------------------------------------------------------------------------


def jpeg(color: str = "red") -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (20, 20), color).save(buffer, "JPEG")
    return buffer.getvalue()


@pytest.fixture
def downloads(monkeypatch):
    """Подмена скачивания: ссылка → байты (или исключение PhotoError)."""
    from apps.catalog import photos

    files: dict[str, bytes] = {}
    calls: list[str] = []

    def fake_download(url: str):
        calls.append(url)
        if url not in files:
            raise photos.PhotoError("the site responded with error 404")
        return photos.to_photo(files[url])

    monkeypatch.setattr(photos, "download", fake_download)
    files["calls"] = calls  # type: ignore[assignment]
    return files


IMAGES_COLUMN = [column.key for column in COLUMNS].index("images")
PHOTO_HEADER = ["slug", "brand", "category", "name_en", "price", "stock_quantity", "images"]


def test_import_downloads_photos_in_order(refs, downloads):
    downloads["https://cdn.example.com/a.jpg"] = jpeg("red")
    downloads["https://cdn.example.com/b.png"] = jpeg("blue")
    cell = "https://cdn.example.com/a.jpg\nhttps://cdn.example.com/b.png"

    report = import_products(
        xlsx(PHOTO_HEADER, ["", "rolex", "watches", "Sub", 100, 1, cell]), commit=True
    )

    assert report.ok, report.errors
    images = list(Product.objects.get(name_en="Sub").images.order_by("position"))
    assert [image.position for image in images] == [0, 1]
    assert images[0].image.name.endswith("rolex-sub-1.jpg")
    assert images[0].width == 20


def test_check_step_does_not_store_photos(refs, downloads):
    downloads["https://cdn.example.com/a.jpg"] = jpeg()

    report = import_products(
        xlsx(
            PHOTO_HEADER, ["", "rolex", "watches", "Sub", 100, 1, "https://cdn.example.com/a.jpg"]
        ),
        commit=False,
    )

    assert report.ok
    assert ProductImage.objects.count() == 0


def test_bad_photo_link_imports_nothing(refs, downloads):
    report = import_products(
        xlsx(
            PHOTO_HEADER,
            ["", "rolex", "watches", "Sub", 100, 1, "https://cdn.example.com/missing.jpg"],
        ),
        commit=True,
    )

    assert [(e.row, e.column) for e in report.errors] == [(2, "images")]
    assert "404" in report.errors[0].message
    assert Product.objects.count() == 0


def test_export_photo_links_round_trip_and_reorder(refs, downloads):
    product = ProductFactory(slug="sub", brand=refs["brand"], category=refs["category"])
    first = ProductImage.objects.create(
        product=product, position=0, image=SimpleUploadedFile("1.jpg", jpeg())
    )
    second = ProductImage.objects.create(
        product=product, position=1, image=SimpleUploadedFile("2.jpg", jpeg())
    )
    exported = export_products(Product.objects.all())

    assert import_products(exported, commit=False).unchanged == 1

    rows = list(load_workbook(io.BytesIO(exported))["Products"].values)
    links = rows[1][rows[0].index("images")].split("\n")
    assert links[0].startswith("http")
    report = import_products(
        xlsx(["slug", "images"], ["sub", "\n".join(reversed(links))]), commit=True
    )

    assert report.rows[0].changes == ["images"]
    assert list(product.images.order_by("position")) == [second, first]
    assert downloads["calls"] == []  # свои фото не скачиваются


def test_photos_replaced_and_copied_from_other_product(refs, downloads):
    donor = ProductFactory(brand=refs["brand"])
    donor_image = ProductImage.objects.create(
        product=donor, image=SimpleUploadedFile("d.jpg", jpeg())
    )
    product = ProductFactory(slug="sub", brand=refs["brand"])
    ProductImage.objects.create(product=product, image=SimpleUploadedFile("old.jpg", jpeg()))
    donor_link = export_products(Product.objects.filter(pk=donor.pk))
    link = list(load_workbook(io.BytesIO(donor_link))["Products"].values)[1][IMAGES_COLUMN]

    report = import_products(xlsx(["slug", "images"], ["sub", link]), commit=True)

    assert report.ok, report.errors
    images = list(product.images.all())
    assert len(images) == 1
    assert images[0].pk != donor_image.pk  # копия файла, а не тот же объект
    assert downloads["calls"] == []


def test_photo_helpers():
    from apps.catalog.photos import PhotoError, check_public_host, direct_link, to_photo

    assert direct_link("https://drive.google.com/file/d/AbC-12/view?usp=sharing") == (
        "https://drive.google.com/uc?export=download&id=AbC-12"
    )
    assert direct_link("https://www.dropbox.com/s/x/watch.jpg?dl=0").endswith("dl=1")
    assert to_photo(jpeg()).extension == "jpg"
    with pytest.raises(PhotoError):
        to_photo(b"<html>not an image</html>")
    with pytest.raises(PhotoError):
        check_public_host("127.0.0.1")
    with pytest.raises(PhotoError):
        check_public_host("192.168.1.10")
