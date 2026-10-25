"""Выгрузки в Excel из админки: общий вид таблицы и ответ с файлом."""

import io
from collections.abc import Iterable, Sequence
from typing import Any

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(content: bytes, name: str) -> HttpResponse:
    """Файл на скачивание: maison-customers-2026-10-25.xlsx."""
    stamp = timezone.localdate().strftime("%Y-%m-%d")
    response = HttpResponse(content, content_type=XLSX)
    response["Content-Disposition"] = f'attachment; filename="{name}-{stamp}.xlsx"'
    return response


def table(
    title: str, headers: Sequence[str], rows: Iterable[Sequence[Any]], widths: Sequence[int]
) -> bytes:
    """Простая таблица: тёмная строка заголовков, закреплена при прокрутке."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
